#!/usr/bin/env python3
"""Deep analysis of a job-application tracker workbook.

Reads an .xlsx with an "Applications" sheet (one row per application) and an
optional "Events" sheet (one row per event), then computes:

  * the stage funnel and reply mix
  * reply / advance rates by segment with Wilson 95% intervals
  * Kaplan-Meier time to first employer reply, treating applications still
    waiting as right-censored instead of as failures
  * Aalen-Johansen cumulative incidence for competing first replies
    (advanced vs. declined)
  * empirical-Bayes (beta-binomial) shrinkage of per-sector rates
  * the chance that a silent application still gets a reply, by days of silence
  * a forecast of replies and screens still to come from the open pipeline

Outputs analysis.json (for the report) and, with --write-xlsx, a copy of the
workbook with "Analysis" and "Open pipeline" sheets added.

Usage:
  python3 analyze_tracker.py TRACKER.xlsx --as-of 2026-10-23 --out-dir out/ \
      --write-xlsx out/tracker_with_analysis.xlsx
"""
import argparse
import datetime as dt
import json
import math
import os
import re
import shutil

import numpy as np
import pandas as pd
from scipy import optimize, special, stats

POSITIVE_STAGES = {"Screen / assessment", "Interview", "Referred / eligible list"}
OPEN_OUTCOMES = {"Open - no decision", "In process", "Stalled - no closure"}

SECTOR = {
    "Specialty insurance": "Insurance / underwriting",
    "Private intelligence": "Private intelligence & investigations",
    "Investigations": "Private intelligence & investigations",
    "Corporate security": "Private intelligence & investigations",
    "Corporate risk": "Private intelligence & investigations",
    "Government - local": "NYC & NY State government",
    "Government - state": "NYC & NY State government",
    "Federal": "Federal government",
    "Government / diplomatic": "Federal government",
    "Government contracting": "Federal government",
    "Media": "Media & communications",
    "Publishing": "Media & communications",
    "Strategic communications": "Media & communications",
    "Media / talent": "Media & communications",
    "Policy research": "Policy & nonprofit",
    "Nonprofit / policy": "Policy & nonprofit",
    "Nonprofit / due diligence": "Policy & nonprofit",
    "Nonprofit / arts": "Policy & nonprofit",
    "Nonprofit / tech": "Policy & nonprofit",
    "Health / policy": "Policy & nonprofit",
    "Financial crime": "Finance, consulting & legal",
    "Finance / research": "Finance, consulting & legal",
    "Finance": "Finance, consulting & legal",
    "Consulting": "Finance, consulting & legal",
    "Economic consulting": "Finance, consulting & legal",
    "Legal / professional services": "Finance, consulting & legal",
    "Research / expert networks": "Finance, consulting & legal",
    "Research": "Finance, consulting & legal",
    "Tech": "Tech, retail & other",
    "Tech / prediction markets": "Tech, retail & other",
    "Corporate / analytics": "Tech, retail & other",
    "Marketing / analytics": "Tech, retail & other",
    "Luxury / retail": "Tech, retail & other",
    "Other": "Tech, retail & other",
}
GOV_PORTALS = {"NYC Jobs portal", "OAG application portal", "State portal"}


# ----------------------------------------------------------------------------- loading

def channel_group(ch):
    if ch == "Employee referral":
        return "Referral link"
    if ch in GOV_PORTALS:
        return "NYC / NY State portal"
    if ch == "USAJOBS":
        return "USAJOBS"
    if ch == "Email application":
        return "Direct email"
    if ch == "LinkedIn":
        return "LinkedIn Easy Apply"
    return "Employer ATS / careers site"


def default_outcome(status, stage):
    s = str(status).lower()
    if stage == "Closed":
        if "withdr" in s:
            return "Withdrew"
        if "filled" in s or "role closed" in s or "hiring complete" in s:
            return "Role filled / closed"
        if "interview" in s or "assessment" in s:
            return "Declined - after screen/interview"
        return "Declined - no interview"
    return "Open - no decision"


def work_mode(role, notes):
    text = f"{role} {notes}".lower()
    if re.search(r"\bremote\b", text) and "not remote" not in text:
        return "Remote"
    if "hybrid" in text:
        return "Hybrid"
    if re.search(r"on[- ]site|in person|fully on site|in-person", text):
        return "On-site"
    return "Not recorded"


SAL_RANGE = re.compile(r"\$(\d{2,3}(?:,\d{3})?)(K)?\s*(?:-|to)\s*\$?(\d{2,3}(?:,\d{3})?)(K)?", re.I)
SAL_ONE = re.compile(r"\$(\d{2,3}),(\d{3})\b|\$(\d{2,3})K\b", re.I)


def salary_mid(notes):
    """Midpoint of the first posted pay range in the notes, in $K (None if absent)."""
    def val(num, k):
        v = float(num.replace(",", ""))
        return v if (k or v < 1000) else v / 1000
    m = SAL_RANGE.search(notes)
    if m:
        hi_k = m.group(4)
        lo = val(m.group(1), m.group(2) or hi_k)
        hi = val(m.group(3), hi_k)
        if 20 <= lo <= 300 and 20 <= hi <= 300:
            return (lo + hi) / 2
    m = SAL_ONE.search(notes)
    if m:
        v = float(m.group(1) + m.group(2)) / 1000 if m.group(1) else float(m.group(3))
        if 20 <= v <= 300:
            return v
    return None


def load(path, as_of):
    apps = pd.read_excel(path, sheet_name="Applications", keep_default_na=False, na_values=[""])
    try:
        events = pd.read_excel(path, sheet_name="Events", keep_default_na=False, na_values=[""])
    except ValueError:
        events = pd.DataFrame(columns=["App ID", "Date", "Event", "Event group"])

    apps["Notes"] = apps["Notes"].fillna("").astype(str)
    if "App ID" not in apps:
        apps = apps.sort_values("Date Applied").reset_index(drop=True)
        apps["App ID"] = [f"A{i:03d}" for i in range(1, len(apps) + 1)]
    defaults = {
        "Sector (grouped)": apps["Track"].map(lambda t: SECTOR.get(t, t)),
        "Channel (grouped)": apps["Channel"].map(channel_group),
        "Insider contact": np.where(apps["Channel"].eq("Employee referral"), "Formal referral", "No insider contact"),
        "Outcome": [default_outcome(s, g) for s, g in zip(apps["Status"], apps["Stage"])],
        "In rate stats": "Yes",
    }
    for col, val in defaults.items():
        if col not in apps:
            apps[col] = val
        else:
            apps[col] = apps[col].fillna(pd.Series(val, index=apps.index) if not np.isscalar(val) else val)
    if "Furthest stage" not in apps:
        apps["Furthest stage"] = np.where(apps["Stage"].eq("Closed"), "Declined / closed, no screen", "No employer reply")
    if "First reply" not in apps:
        apps["First reply"] = np.where(apps["Stage"].eq("Closed"), apps["Last Contact"], pd.NaT)
    apps["First reply"] = pd.to_datetime(apps["First reply"])

    submitted = (events[events["Event"].eq("Application submitted")]
                 .assign(Date=lambda e: pd.to_datetime(e["Date"]))
                 .groupby("App ID")["Date"].min())
    apps["Submitted"] = apps["App ID"].map(submitted).fillna(apps["Date Applied"])
    apps["Submitted"] = pd.to_datetime(apps["Submitted"])

    apps["replied"] = apps["First reply"].notna()
    apps["advanced"] = apps["Furthest stage"].isin(POSITIVE_STAGES)
    apps["interviewed"] = apps["Furthest stage"].eq("Interview")
    apps["declined"] = apps["Outcome"].str.startswith("Declined") | apps["Outcome"].eq("Role filled / closed")
    apps["open"] = apps["Outcome"].isin(OPEN_OUTCOMES)
    # an advanced row with no first-reply date has unknown timing (e.g. interviewed before applying)
    apps["timing_ok"] = ~(apps["advanced"] & ~apps["replied"])
    end = apps["First reply"].where(apps["replied"], pd.Timestamp(as_of))
    apps["duration"] = (end - apps["Submitted"]).dt.days.clip(lower=0)
    apps["reply_cause"] = np.select([apps["replied"] & apps["advanced"], apps["replied"]], ["advanced", "declined"], "none")
    apps["Month"] = apps["Submitted"].dt.to_period("M").astype(str)
    apps["Work mode"] = [work_mode(r, n) for r, n in zip(apps["Role / Req"], apps["Notes"])]
    apps["Salary mid $K"] = apps["Notes"].map(salary_mid)
    apps["Any insider contact"] = np.where(apps["Insider contact"].eq("No insider contact"), "No", "Yes")
    return apps, events


# ----------------------------------------------------------------------------- statistics

def wilson(k, n, z=1.96):
    if n == 0:
        return (None, None)
    p = k / n
    den = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / den
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
    return (max(0.0, centre - half), min(1.0, centre + half))


def kaplan_meier(durations, events):
    """Return event times, S(t), and Greenwood log-log 95% bounds."""
    d = np.asarray(durations, float)
    e = np.asarray(events, bool)
    times = np.unique(d[e])
    s, gw = 1.0, 0.0
    out = []
    for t in times:
        n = (d >= t).sum()
        k = ((d == t) & e).sum()
        s *= 1 - k / n
        if n > k:
            gw += k / (n * (n - k))
        if 0 < s < 1:
            se = math.sqrt(gw) / abs(math.log(s))
            lo, hi = s ** math.exp(1.96 * se), s ** math.exp(-1.96 * se)
        else:
            lo = hi = s
        out.append((float(t), s, lo, hi))
    return out


def surv_at(km, t):
    s = (1.0, 1.0, 1.0)
    for time, sv, lo, hi in km:
        if time <= t:
            s = (sv, lo, hi)
        else:
            break
    return s


def aalen_johansen(durations, causes, cause_names):
    """Cumulative incidence per competing cause; causes == 'none' are censored."""
    d = np.asarray(durations, float)
    c = np.asarray(causes)
    ev = c != "none"
    times = np.unique(d[ev])
    s_prev = 1.0
    cif = {k: 0.0 for k in cause_names}
    out = []
    for t in times:
        n = (d >= t).sum()
        at_t = (d == t) & ev
        total = at_t.sum()
        for k in cause_names:
            cif[k] += s_prev * ((c == k) & at_t).sum() / n
        s_prev *= 1 - total / n
        out.append((float(t), s_prev, dict(cif)))
    return out


def cif_at(aj, t, cause):
    v = 0.0
    for time, _, c in aj:
        if time <= t:
            v = c[cause]
        else:
            break
    return v


def beta_binomial_fit(k, n):
    k, n = np.asarray(k, float), np.asarray(n, float)

    def nll(params):
        a, b = np.exp(params)
        return -np.sum(special.betaln(k + a, n - k + b) - special.betaln(a, b))

    p0 = max(k.sum() / n.sum(), 1e-3)
    res = optimize.minimize(nll, x0=np.log([p0 * 10, (1 - p0) * 10]), method="Nelder-Mead")
    a, b = np.exp(res.x)
    return float(a), float(b)


def poisson_binomial(ps):
    dist = np.array([1.0])
    for p in ps:
        dist = np.convolve(dist, [1 - p, p])
    return dist


def quantile_from_pmf(pmf, q):
    return int(np.searchsorted(np.cumsum(pmf), q))


def pct(x):
    return None if x is None or (isinstance(x, float) and math.isnan(x)) else round(100 * x, 1)


def conditional_probs(km, aj, silent_days):
    """P(any reply) and P(advance) still to come for applications silent for t days."""
    tau = km[-1][0]
    s_tau = surv_at(km, tau)[0]
    adv_tau = cif_at(aj, tau, "advanced")
    p_rep, p_adv = [], []
    for t in silent_days:
        s_t = surv_at(km, t)[0]
        p_rep.append(max(0.0, (s_t - s_tau) / s_t) if s_t > 0 else 0.0)
        p_adv.append(max(0.0, (adv_tau - cif_at(aj, t, "advanced")) / s_t) if s_t > 0 else 0.0)
    return p_rep, p_adv


def forecast_summary(p_rep, p_adv):
    pmf_rep, pmf_adv = poisson_binomial(p_rep), poisson_binomial(p_adv)
    return {
        "expected_replies": round(float(np.sum(p_rep)), 1),
        "replies_80pct": [quantile_from_pmf(pmf_rep, .1), quantile_from_pmf(pmf_rep, .9)],
        "expected_advances": round(float(np.sum(p_adv)), 1),
        "advances_80pct": [quantile_from_pmf(pmf_adv, .1), quantile_from_pmf(pmf_adv, .9)],
        "p_at_least_one_advance": pct(1 - pmf_adv[0]),
    }


# ----------------------------------------------------------------------------- analysis

def segment_table(df, col, km_horizon=30, order=None):
    rows = []
    groups = order or df[col].value_counts().index.tolist()
    for g in groups:
        sub = df[df[col] == g]
        if sub.empty:
            continue
        n, k_rep, k_adv = len(sub), int(sub["replied"].sum()), int(sub["advanced"].sum())
        t = sub[sub["timing_ok"]]
        km = kaplan_meier(t["duration"], t["replied"])
        s30 = surv_at(km, km_horizon)
        mature = sub[(sub["duration"] >= km_horizon) | sub["replied"]]
        k_m30 = int((mature["replied"] & (mature["duration"] <= km_horizon)).sum())
        # segment vs. everyone else, on applications old enough to have been answered
        all_mature = df[((df["duration"] >= km_horizon) | df["replied"]) & df["timing_ok"]]
        inside = all_mature[col] == g
        r30 = all_mature["replied"] & (all_mature["duration"] <= km_horizon)
        p_rep = fisher_vs_rest(r30, inside)
        p_adv = fisher_vs_rest(df["advanced"][df["duration"] >= 14], (df[col] == g)[df["duration"] >= 14])
        rows.append({
            "segment": g, "n": n,
            "replied": k_rep, "reply_rate": pct(k_rep / n), "reply_ci": [pct(x) for x in wilson(k_rep, n)],
            "advanced": k_adv, "advance_rate": pct(k_adv / n), "advance_ci": [pct(x) for x in wilson(k_adv, n)],
            "interviews": int(sub["interviewed"].sum()),
            "km_reply_by_30d": pct(1 - s30[0]),
            "mature_n_30d": len(mature), "mature_replied_30d": k_m30,
            "open": int(sub["open"].sum()),
            "p_reply30_vs_rest": p_rep, "p_advance_vs_rest": p_adv,
        })
    return rows


def fisher_vs_rest(outcome, inside):
    a, b = outcome[inside], outcome[~inside]
    if a.empty or b.empty:
        return None
    _, p = stats.fisher_exact([[int(a.sum()), int((~a).sum())], [int(b.sum()), int((~b).sum())]])
    return round(float(p), 3)


def conditional_table(km_all, aj, horizons):
    tau = km_all[-1][0] if km_all else 0
    s_tau = surv_at(km_all, tau)[0]
    adv_tau = cif_at(aj, tau, "advanced")
    rows = []
    for t in horizons:
        s_t = surv_at(km_all, t)[0]
        rows.append({
            "days_silent": t,
            "p_any_reply": pct((s_t - s_tau) / s_t) if s_t > 0 else None,
            "p_advance": pct((adv_tau - cif_at(aj, t, "advanced")) / s_t) if s_t > 0 else None,
        })
    return rows


def analyse(apps, events, as_of, recent_since):
    stats_df = apps[apps["In rate stats"].eq("Yes")].copy()
    timing = stats_df[stats_df["timing_ok"]]
    n = len(stats_df)

    km_all = kaplan_meier(timing["duration"], timing["replied"])
    aj = aalen_johansen(timing["duration"], timing["reply_cause"], ["advanced", "declined"])

    funnel = [
        ("Applications", n),
        ("Employer replied", int(stats_df["replied"].sum() + (stats_df["advanced"] & ~stats_df["replied"]).sum())),
        ("Advanced (screen, assessment, referral list or interview)", int(stats_df["advanced"].sum())),
        ("Live interview with the hiring team", int(stats_df["interviewed"].sum())),
        ("Offer", 0),
    ]
    funnel_rows = []
    for i, (name, k) in enumerate(funnel):
        prev = funnel[i - 1][1] if i else None
        funnel_rows.append({"stage": name, "count": k,
                            "pct_of_apps": pct(k / n),
                            "conversion_from_prev": pct(k / prev) if prev else None})

    mix = {
        "no_reply": int((~stats_df["replied"] & ~stats_df["advanced"]).sum()),
        "declined_or_closed_without_screen": int((stats_df["replied"] & ~stats_df["advanced"]).sum()),
        "advanced": int(stats_df["advanced"].sum()),
    }

    def days_summary(s):
        s = s.dropna()
        if s.empty:
            return None
        return {"n": int(len(s)), "median": float(s.median()), "p25": float(s.quantile(.25)),
                "p75": float(s.quantile(.75)), "max": float(s.max()),
                "within_14": int((s <= 14).sum()), "within_30": int((s <= 30).sum())}

    pos_days = timing.loc[timing["replied"] & timing["advanced"], "duration"]
    neg_days = timing.loc[timing["replied"] & ~timing["advanced"], "duration"]

    km_points = [{"t": t, "reply_prob": round(1 - s, 4), "lo": round(1 - hi, 4), "hi": round(1 - lo, 4)}
                 for t, s, lo, hi in km_all]
    aj_points = [{"t": t, "advanced": round(c["advanced"], 4), "declined": round(c["declined"], 4)} for t, _, c in aj]
    marks = {}
    for h in (7, 14, 30, 60, 90, 180):
        s, lo, hi = surv_at(km_all, h)
        marks[str(h)] = {"reply_prob": pct(1 - s), "lo": pct(1 - hi), "hi": pct(1 - lo),
                         "advanced": pct(cif_at(aj, h, "advanced")), "declined": pct(cif_at(aj, h, "declined"))}
    max_follow = float(timing["duration"].max())

    # segments
    sector_rows = segment_table(stats_df, "Sector (grouped)")
    channel_rows = segment_table(stats_df, "Channel (grouped)")
    insider_rows = segment_table(stats_df, "Insider contact",
                                 order=["Formal referral", "Insider outreach", "Recruiter-initiated", "No insider contact"])
    insider_any = segment_table(stats_df, "Any insider contact", order=["Yes", "No"])
    month_rows = segment_table(stats_df, "Month", order=sorted(stats_df["Month"].unique()))
    mode_rows = segment_table(stats_df, "Work mode", order=["Remote", "Hybrid", "On-site", "Not recorded"])

    tests = {}
    for label, col in (("advanced", "advanced"), ("replied", "replied")):
        a = stats_df[stats_df["Any insider contact"].eq("Yes")][col]
        b = stats_df[stats_df["Any insider contact"].eq("No")][col]
        table = [[int(a.sum()), int((~a).sum())], [int(b.sum()), int((~b).sum())]]
        odds, p = stats.fisher_exact(table)
        tests[f"insider_vs_none_{label}"] = {"table": table, "odds_ratio": None if math.isinf(odds) else round(float(odds), 2),
                                             "p_value": round(float(p), 4)}

    # empirical Bayes: sector reply-within-30-days among applications with 30+ days of follow-up
    mature = stats_df[(stats_df["duration"] >= 30) | stats_df["replied"]]
    mature = mature[mature["timing_ok"]]
    g = mature.assign(r30=mature["replied"] & (mature["duration"] <= 30)).groupby("Sector (grouped)")["r30"]
    k, nn = g.sum(), g.count()
    a_, b_ = beta_binomial_fit(k.values, nn.values)
    full_pooling = a_ + b_ > 1000  # no between-sector variance beyond binomial noise
    if full_pooling:
        a_, b_ = 1000 * k.sum() / nn.sum(), 1000 * (1 - k.sum() / nn.sum())
    chi2, het_p, _, _ = stats.chi2_contingency(np.column_stack([k.values, nn.values - k.values]))
    eb_rows = []
    for sec in nn.sort_values(ascending=False).index:
        post_a, post_b = k[sec] + a_, nn[sec] - k[sec] + b_
        eb_rows.append({"segment": sec, "n": int(nn[sec]), "replied_30d": int(k[sec]),
                        "raw": pct(k[sec] / nn[sec]), "shrunk": pct(post_a / (post_a + post_b)),
                        "cri": [pct(stats.beta.ppf(.05, post_a, post_b)), pct(stats.beta.ppf(.95, post_a, post_b))]})
    eb = {"prior_alpha": round(a_, 2), "prior_beta": round(b_, 2), "prior_mean": pct(a_ / (a_ + b_)),
          "full_pooling": bool(full_pooling), "heterogeneity_chi2_p": round(float(het_p), 3),
          "pooled_raw": pct(k.sum() / nn.sum()), "rows": eb_rows}

    # salary (exploratory)
    sal = stats_df.dropna(subset=["Salary mid $K"]).copy()
    bins = [0, 60, 80, 100, 1000]
    labels = ["Under $60K", "$60-80K", "$80-100K", "$100K+"]
    sal["Pay band"] = pd.cut(sal["Salary mid $K"], bins=bins, labels=labels, right=False).astype(str)
    salary_rows = segment_table(sal, "Pay band", order=labels)

    # conditional chance of a reply after t days of silence
    cond = conditional_table(km_all, aj, [0, 7, 14, 21, 30, 45, 60, 90, 120])

    # open pipeline forecast: open applications with no reply yet
    pipe = stats_df[stats_df["open"] & ~stats_df["replied"] & stats_df["timing_ok"]].copy()
    pipe["p_reply"], pipe["p_advance"] = conditional_probs(km_all, aj, pipe["duration"])

    def action(row):
        if row["duration"] <= 14:
            return "Wait: still inside the window where most replies arrive"
        if row["duration"] <= 30:
            return "Follow up now if you have a contact; replies are thinning"
        if row["p_advance"] < 0.02:
            return "Treat as no response (mark Silent) unless someone inside can revive it"
        return "Long shot: one polite status check, then let it go"

    pipe["Suggested action"] = pipe.apply(action, axis=1)
    pipe = pipe.sort_values(["p_advance", "p_reply"], ascending=False)
    forecast = forecast_summary(pipe["p_reply"], pipe["p_advance"])
    forecast.update({
        "open_without_reply": int(len(pipe)),
        "young_le_14d": int((pipe["duration"] <= 14).sum()),
        "silent_15_45d": int(((pipe["duration"] > 14) & (pipe["duration"] <= 45)).sum()),
        "silent_gt_45d": int((pipe["duration"] > 45).sum()),
        "live_but_silent_gt_45d": pipe.loc[(pipe["duration"] > 45) & pipe["Stage"].eq("Live"),
                                           ["App ID", "Employer", "duration"]].rename(columns={"duration": "days"}).to_dict("records"),
        "median_age_by_sector": {k: float(v) for k, v in pipe.groupby("Sector (grouped)")["duration"].median().items()},
    })
    # sensitivity: reference curves from recent applications only
    recent = timing[timing["Submitted"] >= pd.Timestamp(recent_since)]
    km_r = kaplan_meier(recent["duration"], recent["replied"])
    aj_r = aalen_johansen(recent["duration"], recent["reply_cause"], ["advanced", "declined"])
    pr, pa = conditional_probs(km_r, aj_r, pipe["duration"])
    forecast["recent_reference"] = {"since": str(recent_since), "n": int(len(recent)), **forecast_summary(pr, pa),
                                    "reply_by_30d": pct(1 - surv_at(km_r, 30)[0]),
                                    "advanced_by_30d": pct(cif_at(aj_r, 30, "advanced"))}
    engaged = stats_df[stats_df["open"] & (stats_df["replied"] | stats_df["advanced"])]

    # weekly volume
    wk = stats_df.assign(week=stats_df["Submitted"].dt.to_period("W-SUN").dt.start_time)
    cat = np.select([wk["advanced"], wk["replied"]], ["Advanced", "Declined / closed"], "No reply yet")
    weekly = (wk.assign(cat=cat).groupby(["week", "cat"]).size().unstack(fill_value=0)
              .reindex(columns=["Advanced", "Declined / closed", "No reply yet"], fill_value=0))
    weekly_rows = [{"week": d.strftime("%Y-%m-%d"), **{c: int(v) for c, v in r.items()}} for d, r in weekly.iterrows()]

    # repeat employers
    emp_key = stats_df["Employer"].str.replace(r"\s*\(.*\)", "", regex=True).str.strip()
    repeats = emp_key.value_counts()
    repeats = repeats[repeats > 1]

    last30 = stats_df[stats_df["Submitted"] >= pd.Timestamp(as_of) - pd.Timedelta(days=30)]
    overview = {
        "as_of": str(as_of), "apps_on_file": int(len(apps)), "apps_in_stats": n,
        "excluded": apps.loc[apps["In rate stats"].ne("Yes"), ["App ID", "Employer"]].to_dict("records"),
        "first_app": stats_df["Submitted"].min().strftime("%Y-%m-%d"),
        "last_app": stats_df["Submitted"].max().strftime("%Y-%m-%d"),
        "apps_last_30d": int(len(last30)),
        "share_since_sep1": pct((stats_df["Submitted"] >= "2026-09-01").mean()),
        "open": int(stats_df["open"].sum()), "replied": int(stats_df["replied"].sum()),
        "advanced": int(stats_df["advanced"].sum()), "interviewed": int(stats_df["interviewed"].sum()),
        "withdrew": int(stats_df["Outcome"].eq("Withdrew").sum()),
        "declined_after_screen": int(stats_df["Outcome"].eq("Declined - after screen/interview").sum()),
        "max_follow_up_days": max_follow,
        "ghost_plateau": pct(km_all[-1][1]),
        "events_logged": int(len(events)),
        "repeat_employers": {k: int(v) for k, v in repeats.items()},
        "salary_parsed": int(len(sal)),
    }

    advanced_list = [{"App ID": r["App ID"], "Employer": r["Employer"], "Role": r["Role / Req"],
                      "Submitted": r["Submitted"].strftime("%Y-%m-%d"), "Furthest stage": r["Furthest stage"],
                      "Outcome": r["Outcome"], "Insider contact": r["Insider contact"],
                      "days_to_first_reply": int(r["duration"]) if r["replied"] else None}
                     for _, r in stats_df[stats_df["advanced"]].sort_values("Submitted").iterrows()]

    return {
        "overview": overview, "funnel": funnel_rows, "mix": mix,
        "timing": {"km": km_points, "aj": aj_points, "marks": marks,
                   "advanced_days": days_summary(pos_days), "declined_days": days_summary(neg_days)},
        "conditional": cond, "forecast": forecast,
        "segments": {"sector": sector_rows, "channel": channel_rows, "insider": insider_rows,
                     "insider_any": insider_any, "month": month_rows, "work_mode": mode_rows, "salary": salary_rows},
        "tests": tests, "empirical_bayes": eb, "weekly": weekly_rows, "advanced_list": advanced_list,
        "pipeline": [{"App ID": r["App ID"], "Employer": r["Employer"], "Role": r["Role / Req"],
                      "Submitted": r["Submitted"].strftime("%Y-%m-%d"), "Days silent": int(r["duration"]),
                      "Stage": r["Stage"], "Insider contact": r["Insider contact"],
                      "P(reply)": pct(r["p_reply"]), "P(advance)": pct(r["p_advance"]),
                      "Suggested action": r["Suggested action"]} for _, r in pipe.iterrows()],
        "engaged": [{"App ID": r["App ID"], "Employer": r["Employer"], "Role": r["Role / Req"],
                     "Furthest stage": r["Furthest stage"], "Outcome": r["Outcome"]} for _, r in engaged.iterrows()],
    }


# ----------------------------------------------------------------------------- workbook output

def write_sheets(src, dst, res):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if os.path.abspath(src) != os.path.abspath(dst):
        shutil.copyfile(src, dst)
    wb = openpyxl.load_workbook(dst)
    for name in ("Analysis", "Open pipeline"):
        if name in wb.sheetnames:
            del wb[name]
    ws = wb.create_sheet("Analysis")
    head_font = Font(name="Arial", size=12, bold=True)
    col_font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    col_fill = PatternFill("solid", fgColor="FF1F3864")
    body = Font(name="Arial", size=10)
    row = 1

    def title(text, note=None):
        nonlocal row
        ws.cell(row, 1, text).font = head_font
        row += 1
        if note:
            c = ws.cell(row, 1, note)
            c.font = Font(name="Arial", size=9, italic=True)
            row += 1

    def table(headers, rows):
        nonlocal row
        for j, h in enumerate(headers, start=1):
            c = ws.cell(row, j, h)
            c.font, c.fill = col_font, col_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        row += 1
        for r in rows:
            for j, v in enumerate(r, start=1):
                c = ws.cell(row, j, v)
                c.font = body
                c.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        row += 1

    ov = res["overview"]
    title(f"Analysis as of {ov['as_of']}",
          f"{ov['apps_in_stats']} of {ov['apps_on_file']} rows are in the statistics (talent-database registrations, a general "
          f"Schedule A request and out-of-period rows are excluded). Applications still waiting are treated as censored, "
          f"not as failures. Rates are shown with Wilson 95% intervals.")
    title("Funnel")
    table(["Stage", "Count", "% of applications", "Conversion from previous stage"],
          [[f["stage"], f["count"], f["pct_of_apps"], f["conversion_from_prev"]] for f in res["funnel"]])

    m = res["timing"]["marks"]
    title("Time to first employer reply (Kaplan-Meier)",
          "Probability an application has had any substantive reply (decline, screen, referral list) by day t.")
    table(["Day", "Reply probability %", "95% low", "95% high", "Advanced by day t %", "Declined by day t %"],
          [[int(k), v["reply_prob"], v["lo"], v["hi"], v["advanced"], v["declined"]] for k, v in m.items()])

    title("If an application has been silent this long, the chance it still produces...")
    table(["Days silent", "Any reply %", "A screen / interview / referral list %"],
          [[c["days_silent"], c["p_any_reply"], c["p_advance"]] for c in res["conditional"]])

    def seg(name, rows_, note=None):
        title(name, note)
        table(["Segment", "Apps", "Replied", "Reply %", "Reply 95% CI", "Advanced", "Advance %", "Advance 95% CI",
               "Interviews", "KM reply by day 30 %", "Still open"],
              [[r["segment"], r["n"], r["replied"], r["reply_rate"], f"{r['reply_ci'][0]}-{r['reply_ci'][1]}",
                r["advanced"], r["advance_rate"], f"{r['advance_ci'][0]}-{r['advance_ci'][1]}", r["interviews"],
                r["km_reply_by_30d"], r["open"]] for r in rows_])

    s = res["segments"]
    seg("By insider contact", s["insider"])
    t = res["tests"]
    title("Insider contact vs. none: Fisher exact tests",
          f"Advanced: odds ratio {t['insider_vs_none_advanced']['odds_ratio']}, p = {t['insider_vs_none_advanced']['p_value']}. "
          f"Any reply: odds ratio {t['insider_vs_none_replied']['odds_ratio']}, p = {t['insider_vs_none_replied']['p_value']}.")
    seg("By sector", s["sector"])
    eb = res["empirical_bayes"]
    title("Sector reply-within-30-days, shrunk toward the overall rate (empirical Bayes)",
          f"Only applications with 30+ days of follow-up. Beta prior fitted across sectors: mean {eb['prior_mean']}%. "
          f"Test for any sector difference: chi-square p = {eb['heterogeneity_chi2_p']}"
          + (" - no difference beyond chance, so the estimates pool to the overall rate." if eb["full_pooling"] else "."))
    table(["Sector", "Apps (30+ days)", "Replied within 30 days", "Raw %", "Shrunk %", "90% credible interval"],
          [[r["segment"], r["n"], r["replied_30d"], r["raw"], r["shrunk"], f"{r['cri'][0]}-{r['cri'][1]}"] for r in eb["rows"]])
    seg("By channel", s["channel"])
    seg("By month applied", s["month"], "Recent months have had less time to hear back; compare the KM column, not raw reply %.")
    seg("By work arrangement (from role and notes)", s["work_mode"])
    seg("By posted pay (exploratory; midpoint of the range in the notes)", s["salary"])

    fc = res["forecast"]
    title("Open pipeline forecast",
          f"{fc['open_without_reply']} open applications have no reply yet. Expected further replies: {fc['expected_replies']} "
          f"(80% range {fc['replies_80pct'][0]}-{fc['replies_80pct'][1]}). Expected screens/interviews/referral lists: "
          f"{fc['expected_advances']} (80% range {fc['advances_80pct'][0]}-{fc['advances_80pct'][1]}); chance of at least one: "
          f"{fc['p_at_least_one_advance']}%. Using only applications since {fc['recent_reference']['since']} as the "
          f"reference: {fc['recent_reference']['expected_replies']} replies and {fc['recent_reference']['expected_advances']} advances.")
    for col, w in zip("ABCDEFGHIJK", (46, 10, 12, 12, 14, 12, 12, 14, 12, 14, 10)):
        ws.column_dimensions[col].width = w

    ps = wb.create_sheet("Open pipeline")
    headers = ["App ID", "Employer", "Role", "Submitted", "Days silent", "Stage", "Insider contact", "P(reply)",
               "P(advance)", "Suggested action"]
    for j, h in enumerate(headers, start=1):
        c = ps.cell(1, j, h + (" %" if h.startswith("P(") else ""))
        c.font, c.fill = col_font, col_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(res["pipeline"], start=2):
        for j, h in enumerate(headers, start=1):
            c = ps.cell(i, j, r[h])
            c.font = body
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate((9, 30, 44, 12, 9, 10, 18, 10, 11, 48), start=1):
        ps.column_dimensions[get_column_letter(j)].width = w
    ps.freeze_panes = "A2"
    ps.auto_filter.ref = f"A1:J{len(res['pipeline']) + 1}"
    wb.save(dst)


def write_datasets(res, out_dir):
    """Flat JSON tables (arrays of row objects) for a dashboard page."""
    os.makedirs(out_dir, exist_ok=True)
    ov, fc, tm = res["overview"], res["forecast"], res["timing"]
    cond = {c["days_silent"]: c for c in res["conditional"]}
    ins = {r["segment"]: r for r in res["segments"]["insider_any"]}
    adv_days, dec_days = tm["advanced_days"] or {}, tm["declined_days"] or {}
    funnel = res["funnel"]
    advanced = res["advanced_list"]
    summary = {
        "as_of": ov["as_of"], "apps_on_file": ov["apps_on_file"], "apps_in_stats": ov["apps_in_stats"],
        "excluded_n": len(ov["excluded"]), "first_app": ov["first_app"], "last_app": ov["last_app"],
        "apps_last_30d": ov["apps_last_30d"], "share_since_sep1": ov["share_since_sep1"],
        "replied": funnel[1]["count"], "reply_pct": funnel[1]["pct_of_apps"],
        "advanced": funnel[2]["count"], "advance_pct": funnel[2]["pct_of_apps"],
        "interviewed": funnel[3]["count"], "interview_pct": funnel[3]["pct_of_apps"], "offers": funnel[4]["count"],
        "open": ov["open"], "withdrew_after_advancing": sum(a["Outcome"] == "Withdrew" for a in advanced),
        "declined_after_advancing": sum(a["Outcome"].startswith("Declined") for a in advanced),
        "still_open_after_advancing": sum(a["Outcome"] in OPEN_OUTCOMES for a in advanced),
        "ghost_pct": ov["ghost_plateau"], "max_follow_up_days": ov["max_follow_up_days"],
        "adv_timed_n": adv_days.get("n"), "adv_median_days": adv_days.get("median"),
        "adv_within_14": adv_days.get("within_14"), "adv_max_days": adv_days.get("max"),
        "dec_n": dec_days.get("n"), "dec_median_days": dec_days.get("median"), "dec_max_days": dec_days.get("max"),
        "reply_by_14": tm["marks"]["14"]["reply_prob"], "reply_by_30": tm["marks"]["30"]["reply_prob"],
        "adv_by_14": tm["marks"]["14"]["advanced"], "adv_by_30": tm["marks"]["30"]["advanced"],
        "p_adv_at_0": cond[0]["p_advance"], "p_adv_after_14": cond[14]["p_advance"], "p_adv_after_21": cond[21]["p_advance"],
        "p_adv_after_45": cond[45]["p_advance"], "p_reply_after_30": cond[30]["p_any_reply"],
        "p_reply_after_60": cond[60]["p_any_reply"], "p_reply_after_90": cond[90]["p_any_reply"],
        "fc_open": fc["open_without_reply"], "fc_replies": fc["expected_replies"],
        "fc_replies_lo": fc["replies_80pct"][0], "fc_replies_hi": fc["replies_80pct"][1],
        "fc_adv": fc["expected_advances"], "fc_adv_lo": fc["advances_80pct"][0], "fc_adv_hi": fc["advances_80pct"][1],
        "fc_p_any_adv": fc["p_at_least_one_advance"], "fc_recent_since": fc["recent_reference"]["since"],
        "fc_recent_n": fc["recent_reference"]["n"], "fc_recent_replies": fc["recent_reference"]["expected_replies"],
        "fc_recent_adv": fc["recent_reference"]["expected_advances"],
        "young_le_14d": fc["young_le_14d"], "silent_15_45d": fc["silent_15_45d"], "silent_gt_45d": fc["silent_gt_45d"],
        "live_but_silent_gt_45d": len(fc["live_but_silent_gt_45d"]),
        "sector_het_p": res["empirical_bayes"]["heterogeneity_chi2_p"], "sector_pooled_30d": res["empirical_bayes"]["pooled_raw"],
        "insider_n": ins["Yes"]["n"], "insider_adv": ins["Yes"]["advanced"], "insider_interviews": ins["Yes"]["interviews"],
        "insider_open": ins["Yes"]["open"], "none_n": ins["No"]["n"], "none_adv": ins["No"]["advanced"],
        "insider_adv_p": round(res["tests"]["insider_vs_none_advanced"]["p_value"], 3), "events_logged": ov["events_logged"],
    }
    tables = {"summary": [summary]}
    tables["funnel"] = [{"step": i + 1, "stage": f["stage"], "count": f["count"], "pct_of_apps": f["pct_of_apps"],
                         "conversion_pct": f["conversion_from_prev"],
                         "lost": (funnel[i - 1]["count"] - f["count"]) if i else None,
                         "lost_pct": round(100 - f["conversion_from_prev"], 1) if i else None}
                        for i, f in enumerate(funnel)]
    aj = {p["t"]: p for p in tm["aj"]}
    curve = [{"day": 0, "any_reply": 0.0, "any_lo": 0.0, "any_hi": 0.0, "advanced": 0.0, "declined": 0.0}]
    for k in tm["km"]:
        a = aj.get(k["t"], {"advanced": curve[-1]["advanced"] / 100, "declined": curve[-1]["declined"] / 100})
        curve.append({"day": int(k["t"]), "any_reply": pct(k["reply_prob"]), "any_lo": pct(k["lo"]), "any_hi": pct(k["hi"]),
                      "advanced": pct(a["advanced"]), "declined": pct(a["declined"])})
    if curve[-1]["day"] < ov["max_follow_up_days"]:
        curve.append({**curve[-1], "day": int(ov["max_follow_up_days"])})
    tables["reply_curve"] = curve
    tables["silence"] = res["conditional"]
    seg_rows = []
    labels = {"sector": "Sector", "channel": "How you applied", "insider": "Insider contact", "month": "Month applied"}
    for key, label in labels.items():
        for i, r in enumerate(res["segments"][key]):
            seg_rows.append({"seg_key": f"{key}-{i + 1:02d}", "breakdown": label, "segment": r["segment"], "n": r["n"],
                             "replied": r["replied"], "reply_pct": r["reply_rate"], "reply_lo": r["reply_ci"][0],
                             "reply_hi": r["reply_ci"][1], "advanced": r["advanced"], "advance_pct": r["advance_rate"],
                             "advance_lo": r["advance_ci"][0], "advance_hi": r["advance_ci"][1],
                             "interviews": r["interviews"], "reply_by_30_km": r["km_reply_by_30d"], "open": r["open"],
                             "mature_n": r["mature_n_30d"], "mature_replied_30": r["mature_replied_30d"],
                             "p_reply30_vs_rest": r["p_reply30_vs_rest"], "p_advance_vs_rest": r["p_advance_vs_rest"]})
    tables["segments"] = seg_rows
    tables["weekly"] = [{"week": w["week"], "advanced": w["Advanced"], "declined": w["Declined / closed"],
                         "no_reply": w["No reply yet"], "total": w["Advanced"] + w["Declined / closed"] + w["No reply yet"]}
                        for w in res["weekly"]]
    tables["pipeline"] = [{"app_id": p["App ID"], "employer": p["Employer"], "role": p["Role"], "submitted": p["Submitted"],
                           "days_silent": p["Days silent"], "stage": p["Stage"], "insider": p["Insider contact"],
                           "p_reply": p["P(reply)"], "p_advance": p["P(advance)"], "action": p["Suggested action"]}
                          for p in res["pipeline"]]
    tables["advanced"] = [{"app_id": a["App ID"], "employer": a["Employer"], "role": a["Role"], "submitted": a["Submitted"],
                           "stage_reached": a["Furthest stage"], "outcome": a["Outcome"], "insider": a["Insider contact"],
                           "days_to_first_reply": a["days_to_first_reply"]} for a in advanced]
    for name, rows in tables.items():
        with open(os.path.join(out_dir, f"{name}.json"), "w") as f:
            json.dump(rows, f, indent=1)
    return list(tables)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("tracker")
    ap.add_argument("--as-of", default=dt.date.today().isoformat())
    ap.add_argument("--out-dir", default=".")
    ap.add_argument("--write-xlsx", help="write a copy of the tracker with Analysis and Open pipeline sheets")
    ap.add_argument("--dataset-dir", help="also write flat JSON tables for a dashboard page")
    ap.add_argument("--recent-since", default=None,
                    help="start date of the recent-only reference used for the forecast sensitivity (default: 150 days before --as-of)")
    args = ap.parse_args()

    as_of = dt.date.fromisoformat(args.as_of)
    apps, events = load(args.tracker, as_of)
    recent_since = dt.date.fromisoformat(args.recent_since) if args.recent_since else as_of - dt.timedelta(days=150)
    res = analyse(apps, events, as_of, recent_since)
    os.makedirs(args.out_dir, exist_ok=True)
    with open(os.path.join(args.out_dir, "analysis.json"), "w") as f:
        json.dump(res, f, indent=1, default=str)
    if args.write_xlsx:
        write_sheets(args.tracker, args.write_xlsx, res)
    if args.dataset_dir:
        write_datasets(res, args.dataset_dir)
    ov, fc = res["overview"], res["forecast"]
    print(f"{ov['apps_in_stats']} applications analysed; {ov['replied']} replies, {ov['advanced']} advanced, "
          f"{ov['interviewed']} interviews; forecast {fc['expected_replies']} more replies, {fc['expected_advances']} advances")


if __name__ == "__main__":
    main()
